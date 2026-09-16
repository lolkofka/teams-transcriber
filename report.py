"""
Lesson report: transcript + slides merged by time into a DOCX or PDF, from a lesson folder
(<folder>/transcript.txt header, transcript.jsonl, slides.jsonl, slides/*.png).

    python report.py "<lesson folder>" [docx|pdf|both]
"""

import datetime as dt
import json
import os
import sys


def _load(folder):
    header = {}
    txt = os.path.join(folder, "transcript.txt")
    if os.path.exists(txt):
        with open(txt, encoding="utf-8") as f:
            for line in f:
                if not line.startswith("#"):
                    break
                if ": " in line:
                    k, v = line[1:].split(": ", 1)
                    header[k.strip()] = v.strip()
    items = []
    p = os.path.join(folder, "transcript.jsonl")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            for line in f:
                try:
                    d = json.loads(line)
                except Exception:
                    continue
                if d.get("text"):
                    items.append({"kind": "text", "t": d.get("ts") or 0, "wall": d.get("wall", ""), "speaker": d.get("speaker", ""),
                                  "source": d.get("source", "teams"), "text": d["text"]})
    p = os.path.join(folder, "slides.jsonl")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            for line in f:
                try:
                    d = json.loads(line)
                except Exception:
                    continue
                path = os.path.join(folder, d.get("file", ""))
                if os.path.exists(path):
                    items.append({"kind": "slide", "t": d.get("t") or 0, "wall": d.get("wall", ""), "path": path, "presenter": d.get("presenter") or ""})
    items.sort(key=lambda x: x["t"])
    return header, items


def _who(speaker, source):
    if source == "mic" or speaker == "Вы":
        return "Вы"
    return (speaker or "").replace("SPEAKER_", "Спикер ")


def _title(header, folder):
    subject = header.get("Пара") or os.path.basename(os.path.dirname(folder.rstrip("\\/"))) or "Запись"
    when = header.get("Дата") or os.path.basename(folder.rstrip("\\/"))
    return subject, when


def _analysis(folder):
    p = os.path.join(folder, "analysis.json")
    if os.path.exists(p):
        try:
            with open(p, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return None


# ------------------------------------------------------------------------------------------------ DOCX
def build_docx(folder, out=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    header, items = _load(folder)
    subject, when = _title(header, folder)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    doc.add_heading(subject, level=1)
    meta = [f"Дата: {when}"]
    if header.get("Преподаватель"):
        meta.append(f"Преподаватель: {header['Преподаватель']}")
    doc.add_paragraph(" · ".join(meta))
    an = _analysis(folder)
    if an and (an.get("topic") or an.get("task")):
        doc.add_heading("Кратко", level=2)
        if an.get("topic"):
            doc.add_paragraph(f"Тема: {an['topic']}")
        if an.get("task"):
            doc.add_paragraph(f"Задание: {an['task']}" + (f" (срок: {an['deadline']})" if an.get("deadline") else ""))
        for pnt in an.get("points") or []:
            doc.add_paragraph(pnt, style="List Bullet")
    doc.add_heading("Ход пары", level=2)
    n_slides = 0
    for it in items:
        if it["kind"] == "slide":
            n_slides += 1
            doc.add_picture(it["path"], width=Inches(6.0))
            cap = doc.add_paragraph(f"Слайд {n_slides} · {it['wall']}" + (f" · показывает {it['presenter']}" if it["presenter"] else ""))
            cap.runs[0].font.size = Pt(9); cap.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        else:
            p = doc.add_paragraph()
            r = p.add_run(f"[{it['wall']}] "); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            who = _who(it["speaker"], it["source"])
            r = p.add_run(f"{who}: "); r.bold = True
            if who == "Вы":
                r.font.color.rgb = RGBColor(0x0B, 0x7A, 0x75)
            p.add_run(it["text"])
    out = out or os.path.join(folder, "конспект.docx")
    doc.save(out)
    return out


# ------------------------------------------------------------------------------------------------ PDF
def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for name, path in (("SegoeUI", r"C:\Windows\Fonts\segoeui.ttf"), ("Arial", r"C:\Windows\Fonts\arial.ttf"), ("DejaVuSans", r"C:\Windows\Fonts\DejaVuSans.ttf")):
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                bold = path.replace(".ttf", "b.ttf") if name == "SegoeUI" else path.replace("arial.ttf", "arialbd.ttf")
                if os.path.exists(bold):
                    pdfmetrics.registerFont(TTFont(name + "-Bold", bold))
                    return name, name + "-Bold"
                return name, name
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold"


def build_pdf(folder, out=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, ListFlowable, ListItem
    from reportlab.lib import colors
    from xml.sax.saxutils import escape
    header, items = _load(folder)
    subject, when = _title(header, folder)
    font, bold = _pdf_font()
    out = out or os.path.join(folder, "конспект.pdf")
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=subject)
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceAfter=6)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=10, leading=13.5)
    meta = ParagraphStyle("meta", fontName=font, fontSize=9, leading=12, textColor=colors.grey)
    story = [Paragraph(escape(subject), h1)]
    mline = [f"Дата: {when}"] + [f"Преподаватель: {header['Преподаватель']}"] if header.get("Преподаватель") else [f"Дата: {when}"]
    story.append(Paragraph(escape(" · ".join(mline)), meta))
    an = _analysis(folder)
    if an and (an.get("topic") or an.get("task")):
        story.append(Paragraph("Кратко", h2))
        if an.get("topic"):
            story.append(Paragraph("<b>Тема:</b> " + escape(an["topic"]), body))
        if an.get("task"):
            story.append(Paragraph("<b>Задание:</b> " + escape(an["task"]) + (f" (срок: {escape(an['deadline'])})" if an.get("deadline") else ""), body))
        pts = [ListItem(Paragraph(escape(p), body)) for p in (an.get("points") or [])]
        if pts:
            story.append(ListFlowable(pts, bulletType="bullet", leftIndent=12))
    story.append(Paragraph("Ход пары", h2))
    avail_w = A4[0] - 36 * mm
    n_slides = 0
    for it in items:
        if it["kind"] == "slide":
            n_slides += 1
            from PIL import Image as PILImage
            with PILImage.open(it["path"]) as im:
                w, h = im.size
            scale = min(1.0, avail_w / w)
            story.append(Spacer(1, 4))
            story.append(RLImage(it["path"], width=w * scale, height=h * scale))
            story.append(Paragraph(escape(f"Слайд {n_slides} · {it['wall']}" + (f" · показывает {it['presenter']}" if it["presenter"] else "")), meta))
            story.append(Spacer(1, 4))
        else:
            who = _who(it["speaker"], it["source"])
            color = "#0b7a75" if who == "Вы" else "#000000"
            story.append(Paragraph(f'<font color="grey" size="8">[{escape(it["wall"])}]</font> <b><font color="{color}">{escape(who)}:</font></b> {escape(it["text"])}', body))
    doc.build(story)
    return out


def build(folder, fmt="docx"):
    if fmt == "pdf":
        return build_pdf(folder)
    if fmt == "both":
        return build_docx(folder), build_pdf(folder)
    return build_docx(folder)


if __name__ == "__main__":
    folder = sys.argv[1]
    fmt = sys.argv[2] if len(sys.argv) > 2 else "both"
    print(build(folder, fmt))
